from docx import Document

DOC_PATH = "2025年度中国人民大学科研标兵评审表.docx"
OUTPUT_PATH = "科研标兵评审表_填好表五.docx"

# ---------- 1. 载入 ----------
doc = Document(DOC_PATH)

# ---------- 2. 找到“表五” ----------
def find_table_by_keyword(document, keyword: str):
    """
    返回第一张包含指定关键字的表格；若找不到则抛异常
    """
    for tbl in document.tables:
        whole_text = "".join(
            cell.text for row in tbl.rows for cell in row.cells
        ).replace(" ", "").replace("\n", "")
        if keyword in whole_text:
            return tbl
    raise ValueError(f"未找到包含关键字“{keyword}”的表格")

table5 = find_table_by_keyword(doc, "表五")

# ---------- 3. 准备要写入的数据 ----------
# 假设“表五”列布局：
# | 序号 | 项目名称 | 来源 | 经费(万元) | 起止时间 | 本人排名 |
rows_to_insert = [
    ("1", "国产开源 LLM 可信验证平台", "国家自然科学基金", "320", "2024.1–2026.12", "1/5"),
    ("2", "大模型加速推理编译器", "企业横向",          "150", "2023.5–2024.10", "1/4"),
]

# ---------- 4. 写入 ----------
# 先看模板是否只剩表头一行；如果空行不够就 add_row
rows_in_tpl = len(table5.rows)
rows_needed = 1 + len(rows_to_insert)        # 1 行表头 + N 行数据
for _ in range(rows_needed - rows_in_tpl):
    table5.add_row()

# 假设第 0 行是表头，从第 1 行开始写
start_idx = 1
for i, data_row in enumerate(rows_to_insert, start=start_idx):
    for j, cell_val in enumerate(data_row):
        table5.cell(i, j).text = str(cell_val)

# ---------- 5. 保存 ----------
doc.save(OUTPUT_PATH)
print(f"✅ 表五已写入完成 → {OUTPUT_PATH}")
